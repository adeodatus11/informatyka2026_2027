#!/usr/bin/env python3
"""Pliki testowe dla sprawdzarki docxCheck (tests/sims-docx.test.mjs).

Z plików startowych lekcji 28–33 tworzy „rozwiązane” dokumenty z prawdziwymi strukturami Worda:
style, nagłówki/stopki z polem PAGE (pole złożone, fldSimple i pole w kontrolce zawartości),
spis treści TOC z wynikiem i zakładkami _Toc, podpisy SEQ, spisy ilustracji TOC \\c, odsyłacz REF,
sekcje (orient=landscape, kolumny, podział Ciągły, rozłączone nagłówki), komentarze, w:ins/w:del.
Oraz przypadki błędne: nieaktualny spis, ręczny numer strony, udawany nagłówek, uszkodzony ZIP.

Uruchom: python3 scripts/build-docx-fixtures.py   (python-docx ≥ 1.2, lxml)
Wynik: tests/fixtures/docx/*.docx — małe pliki (bez miniatury i stylesWithEffects).
"""
from pathlib import Path
import copy, io, re, zipfile
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.text.run import Run

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / 'public/materials/word'
OUT = ROOT / 'tests/fixtures/docx'
OUT.mkdir(parents=True, exist_ok=True)
W = nsdecls('w')
WR = nsdecls('w', 'r')
DATE = '2026-10-20T10:00:00Z'
_rev = [100]


def rid():
    _rev[0] += 1
    return str(_rev[0])


def start(n):
    return Document(str(SRC / f'lesson-{n}-start.docx'))


def para(doc, text, exact=True):
    for p in doc.paragraphs:
        t = p.text.strip()
        if (t == text) if exact else t.startswith(text):
            return p
    raise KeyError(text)


def set_style(p, name):
    p.style = p.part.document.styles[name]
    for r in p.runs:
        r.bold = None
        r.font.size = None


def esc(s):
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


def run_xml(text, rpr=''):
    return f'<w:r>{rpr}<w:t xml:space="preserve">{esc(text)}</w:t></w:r>'


def field_runs(instr, result, rpr=''):
    """Pole złożone: begin / instrText / separate / wynik / end — tak zapisuje je Word."""
    return (f'<w:r>{rpr}<w:fldChar w:fldCharType="begin"/></w:r>'
            f'<w:r>{rpr}<w:instrText xml:space="preserve"> {esc(instr)} </w:instrText></w:r>'
            f'<w:r>{rpr}<w:fldChar w:fldCharType="separate"/></w:r>'
            f'<w:r>{rpr}<w:t>{esc(result)}</w:t></w:r>'
            f'<w:r>{rpr}<w:fldChar w:fldCharType="end"/></w:r>')


def append_xml(p, xml):
    frag = parse_xml(f'<w:p {WR}>{xml}</w:p>')
    for child in list(frag):
        p._p.append(child)


def clear_runs(p):
    for child in list(p._p):
        if child.tag != qn('w:pPr'):
            p._p.remove(child)


def new_para_after(el, style=None, xml=''):
    ppr = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ''
    p = parse_xml(f'<w:p {WR}>{ppr}{xml}</w:p>')
    el.addnext(p)
    return p


def new_para_before(el, style=None, xml=''):
    ppr = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ''
    p = parse_xml(f'<w:p {WR}>{ppr}{xml}</w:p>')
    el.addprevious(p)
    return p


def bookmark(p_el, name, bid):
    s = parse_xml(f'<w:bookmarkStart {W} w:id="{bid}" w:name="{name}"/>')
    e = parse_xml(f'<w:bookmarkEnd {W} w:id="{bid}"/>')
    ppr = p_el.find(qn('w:pPr'))
    (ppr.addnext if ppr is not None else (lambda x: p_el.insert(0, x)))(s)
    p_el.append(e)


def ensure_style(doc, sid, name, kind='paragraph', extra=''):
    st = doc.styles.element
    if st.find(f'.//w:style[@w:styleId="{sid}"]', namespaces={'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}) is not None:
        return
    st.append(parse_xml(f'<w:style {W} w:type="{kind}" w:styleId="{sid}"><w:name w:val="{name}"/>'
                        f'<w:basedOn w:val="{"Normal" if kind == "paragraph" else "DefaultParagraphFont"}"/>'
                        f'<w:uiPriority w:val="39"/><w:unhideWhenUsed/>{extra}</w:style>'))


def toc_styles(doc):
    ensure_style(doc, 'TOC1', 'toc 1', extra='<w:pPr><w:spacing w:after="100"/></w:pPr>')
    ensure_style(doc, 'TOC2', 'toc 2', extra='<w:pPr><w:spacing w:after="100"/><w:ind w:left="220"/></w:pPr>')
    ensure_style(doc, 'TableofFigures', 'table of figures', extra='<w:pPr><w:spacing w:after="0"/></w:pPr>')
    ensure_style(doc, 'Hyperlink', 'Hyperlink', kind='character', extra='<w:rPr><w:color w:val="0563C1"/><w:u w:val="single"/></w:rPr>')


def toc_paragraphs(entries, instr, style_of, heading=None, sdt=True):
    """Spis jako pole TOC z wynikiem: wpisy z hiperłączami do zakładek i polami PAGEREF."""
    rows = []
    if heading:
        rows.append(f'<w:p><w:pPr><w:pStyle w:val="TOCHeading"/></w:pPr>{run_xml(heading)}</w:p>')
    nop = '<w:rPr><w:noProof/><w:webHidden/></w:rPr>'
    for i, (text, level, bm, page) in enumerate(entries):
        begin = ''
        if i == 0:
            begin = (f'<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> {esc(instr)} </w:instrText></w:r>'
                     f'<w:r><w:fldChar w:fldCharType="separate"/></w:r>')
        pageref = field_runs('PAGEREF ' + bm + ' \\h', str(page), nop)
        rows.append(f'<w:p><w:pPr><w:pStyle w:val="{style_of(level)}"/><w:tabs><w:tab w:val="right" w:leader="dot" w:pos="9350"/></w:tabs></w:pPr>{begin}'
                    f'<w:hyperlink w:anchor="{bm}" w:history="1"><w:r><w:rPr><w:rStyle w:val="Hyperlink"/><w:noProof/></w:rPr><w:t xml:space="preserve">{esc(text)}</w:t></w:r>'
                    f'<w:r>{nop}<w:tab/></w:r>{pageref}</w:hyperlink></w:p>')
    rows.append('<w:p><w:r><w:rPr><w:b/><w:noProof/></w:rPr><w:fldChar w:fldCharType="end"/></w:r></w:p>')
    inner = ''.join(rows)
    if sdt:
        return [parse_xml(f'<w:sdt {WR}><w:sdtPr><w:id w:val="{rid()}"/><w:docPartObj><w:docPartGallery w:val="Table of Contents"/><w:docPartUnique/></w:docPartObj></w:sdtPr>'
                          f'<w:sdtEndPr/><w:sdtContent>{inner}</w:sdtContent></w:sdt>')]
    return list(parse_xml(f'<w:body {WR}>{inner}</w:body>'))


def heading_entries(doc, levels=(1, 2), start_page=2):
    """Zakładki _Toc na nagłówkach + lista wpisów (tekst, poziom, zakładka, strona)."""
    out = []
    page = start_page
    for i, p in enumerate(doc.paragraphs):
        name = p.style.name.lower()
        m = re.match(r'heading (\d)', name)
        if not m or int(m.group(1)) not in levels or not p.text.strip():
            continue
        bm = f'_Toc{2000000 + i}'
        bookmark(p._p, bm, rid())
        out.append((p.text.strip(), int(m.group(1)), bm, page))
        page += 1 if int(m.group(1)) == 1 else 0
    return out


def caption_xml(label, num, desc, bm=None):
    inner = run_xml(label + ' ') + field_runs(f'SEQ {label} \\* ARABIC', str(num), '<w:rPr><w:noProof/></w:rPr>')
    body = inner + run_xml(desc)
    if bm:
        b = rid()
        body = f'<w:bookmarkStart w:id="{b}" w:name="{bm}"/>{body}<w:bookmarkEnd w:id="{b}"/>'
    return body


def header_text(section, text, linked=False):
    hdr = section.header
    hdr.is_linked_to_previous = linked
    if not linked:
        p = hdr.paragraphs[0]
        p.text = text
        p.style = section.part.document.styles['Header'] if hasattr(section.part, 'document') else p.style


def footer_page(section, prefix='', mode='complex'):
    ftr = section.footer
    ftr.is_linked_to_previous = False
    p = ftr.paragraphs[0]
    clear_runs(p)
    if prefix:
        append_xml(p, run_xml(prefix))
    if mode == 'complex':
        append_xml(p, field_runs('PAGE', '1'))
    elif mode == 'simple':
        append_xml(p, f'<w:fldSimple w:instr=" PAGE   \\* MERGEFORMAT "><w:r><w:t>1</w:t></w:r></w:fldSimple>')
    elif mode == 'sdt':  # Wstawianie → Numer strony → Dół strony (Word wstawia pole w kontrolce zawartości)
        page_runs = field_runs('PAGE   \\* MERGEFORMAT', '1')
        sdt = parse_xml(f'<w:sdt {WR}><w:sdtPr><w:id w:val="{rid()}"/><w:docPartObj><w:docPartGallery w:val="Page Numbers (Bottom of Page)"/><w:docPartUnique/></w:docPartObj></w:sdtPr>'
                        f'<w:sdtContent><w:p><w:pPr><w:pStyle w:val="Footer"/><w:jc w:val="center"/></w:pPr>{page_runs}</w:p></w:sdtContent></w:sdt>')
        p._p.addprevious(sdt)
        p._p.getparent().remove(p._p)
    elif mode == 'manual':
        append_xml(p, run_xml('1'))


def sect_xml(landscape=False, cols=1, kind=None):
    w, h = (16838, 11906) if landscape else (11906, 16838)
    orient = ' w:orient="landscape"' if landscape else ''
    t = f'<w:type w:val="{kind}"/>' if kind else ''
    c = f'<w:cols w:num="{cols}" w:space="708"/>' if cols > 1 else '<w:cols w:space="720"/>'
    return (f'<w:sectPr {WR}>{t}<w:pgSz w:w="{w}" w:h="{h}"{orient}/>'
            f'<w:pgMar w:top="1134" w:right="1247" w:bottom="1134" w:left="1247" w:header="720" w:footer="720" w:gutter="0"/>{c}<w:docGrid w:linePitch="360"/></w:sectPr>')


def put_sect(p_el, xml):
    ppr = p_el.find(qn('w:pPr'))
    if ppr is None:
        ppr = OxmlElement('w:pPr')
        p_el.insert(0, ppr)
    ppr.append(parse_xml(xml))


def save(doc, name, stored=False, rename_styles=None):
    buf = io.BytesIO()
    doc.save(buf)
    data = slim(buf.getvalue(), stored, rename_styles)
    (OUT / name).write_bytes(data)
    print(f'{name}: {len(data)} B')
    return data


DROP = ('docProps/thumbnail.jpeg', 'word/stylesWithEffects.xml', 'customXml/')


def slim(data, stored=False, rename_styles=None):
    """Usuwa części nieistotne dla testów (miniatura, stylesWithEffects, customXml), opcjonalnie zmienia nazwy stylów."""
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    zout = zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        if item.filename.startswith(DROP):
            continue
        b = zin.read(item.filename)
        if item.filename.endswith('.rels') or item.filename == '[Content_Types].xml':
            s = b.decode('utf8')
            s = re.sub(r'<Relationship [^>]*Target="(?:\.\./)?(?:customXml/[^"]*|stylesWithEffects\.xml|docProps/thumbnail\.jpeg)"[^>]*/>', '', s)
            s = re.sub(r'<Override [^>]*PartName="/(?:customXml/[^"]*|word/stylesWithEffects\.xml)"[^>]*/>', '', s)
            b = s.encode('utf8')
        if rename_styles and item.filename.startswith('word/') and item.filename.endswith('.xml'):
            s = b.decode('utf8')
            for old_id, new_id, new_name in rename_styles:
                s = re.sub(rf'w:val="{old_id}"', f'w:val="{new_id}"', s)
                s = re.sub(rf'w:styleId="{old_id}"', f'w:styleId="{new_id}"', s)
                if item.filename == 'word/styles.xml':
                    s = re.sub(rf'(<w:style [^>]*w:styleId="{new_id}"[^>]*><w:name w:val=")[^"]*"', rf'\g<1>{new_name}"', s)
            b = s.encode('utf8')
        # stored=True: metoda 0 dla wszystkich części poza dużym styles.xml (archiwum mieszane 0 + 8)
        method = zipfile.ZIP_STORED if stored and item.filename != 'word/styles.xml' else zipfile.ZIP_DEFLATED
        zout.writestr(item.filename, b, compress_type=method)
    zout.close()
    return out.getvalue()


# ——— Lekcja 28: style, nagłówek, stopka z polem PAGE ———
H1_28 = ['Po co zakładamy klub', 'Zasady współpracy', 'Jak sprawdzimy efekty']
H2_28 = ['Pierwsze spotkanie', 'Role w zespole', 'Co przygotować']


def lesson28(variant='solved'):
    d = start(28)
    for t in H1_28:
        if variant == 'fake' and t == 'Jak sprawdzimy efekty':
            continue  # pozostaje pogrubiony Normalny — „udawany nagłówek”
        set_style(para(d, t), 'Heading 1')
    for t in H2_28:
        set_style(para(d, t), 'Heading 2')
    d.styles['Heading 1'].font.size = Pt(16)
    d.styles['Heading 1'].font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    d.styles['Normal'].font.size = Pt(11)
    sec = d.sections[0]
    header_text(sec, 'Szkolny klub gier')
    if variant == 'manual':
        footer_page(sec, 'Poradnik | strona ', 'manual')
    else:
        footer_page(sec, 'Poradnik | strona ', 'simple' if variant == 'pl' else 'complex')
    return d


# ——— Lekcja 29: spis treści ———
def lesson29(variant='solved'):
    d = start(29)
    toc_styles(d)
    first = para(d, 'Cel wydarzenia')
    if variant == 'stale':
        entries = heading_entries(d)          # spis wstawiony przed zmianami…
    first.runs[0].text = 'Cel wydarzenia — wersja szkolna'
    last = d.paragraphs[-1]
    h = d.add_paragraph('Kontakt do organizatora', 'Heading 1')
    d.add_paragraph('Pytania dotyczące turnieju zgłaszamy opiekunowi klubu w szkole.')
    # Ctrl+Enter przed drugim rozdziałem
    new_para_before(para(d, 'Przygotowanie sali')._p, xml='<w:r><w:br w:type="page"/></w:r>')
    if variant != 'stale':
        entries = heading_entries(d)          # …albo zaktualizowany po zmianach
    holder = para(d, 'Miejsce na automatyczny spis treści')
    for el in toc_paragraphs(entries, 'TOC \\o "1-3" \\h \\z \\u', lambda lv: f'TOC{lv}', heading='Spis treści'):
        holder._p.addprevious(el)
    holder._p.getparent().remove(holder._p)
    return d


# ——— Lekcja 30: podpisy SEQ i spisy ilustracji ———
def lesson30():
    d = start(30)
    toc_styles(d)
    body = d.element.body
    pics = [p for p in d.paragraphs if p._p.findall('.//' + qn('w:drawing'))]
    descs = [('Ilustracja', 1, '. Plan strefy nauki — wariant szkolny'), ('Ilustracja', 2, '. Etapy pracy: zaplanuj, wykonaj, sprawdź'),
             ('Wykres', 1, '. Średni czas jednej wizyty w kolejnych dniach'), ('Wykres', 2, '. Potrzeby uczestników według liczby wskazań')]
    caps = {}
    for p, (label, num, desc) in zip(pics, descs):
        bm = f'_Toc{3000000 + len(caps)}'
        new_para_after(p._p, 'Caption', caption_xml(label, num, desc, bm))
        caps.setdefault(label, []).append((f'{label} {num}{desc}', bm))
    tables = body.findall(qn('w:tbl'))
    for i, (tbl, desc) in enumerate(zip(tables, ['. Wyposażenie strefy nauki', '. Dyżury i godziny otwarcia'])):
        bm = f'_Toc{3100000 + i}'
        b = rid()
        cap = caption_xml('Tabela', i + 1, desc, bm)
        if i == 0:  # zakładka _Ref pod odsyłacz „zob. Tabela 1”
            cap = f'<w:bookmarkStart w:id="{b}" w:name="_Ref3100"/>{cap}<w:bookmarkEnd w:id="{b}"/>'
        new_para_before(tbl, 'Caption', cap)
        caps.setdefault('Tabela', []).append((f'Tabela {i + 1}{desc}', bm))
    ref_par = para(d, 'Obie poniższe tabele', exact=False)
    append_xml(ref_par, run_xml(' Zestawienie wyposażenia: zob. ') + f'<w:fldSimple w:instr=" REF _Ref3100 \\h "><w:r><w:t>Tabela 1</w:t></w:r></w:fldSimple>' + run_xml('.'))
    holder = para(d, 'Tutaj wstaw trzy osobne spisy', exact=False)
    anchor = holder._p
    for label, title in [('Ilustracja', 'Spis ilustracji'), ('Tabela', 'Spis tabel'), ('Wykres', 'Spis wykresów')]:
        anchor = new_para_after(anchor, 'Heading2', run_xml(title))
        entries = [(text, 1, bm, 2 + k) for k, (text, bm) in enumerate(caps[label])]
        for el in toc_paragraphs(entries, f'TOC \\h \\z \\c "{label}"', lambda lv: 'TableofFigures', sdt=False):
            anchor.addnext(el)
            anchor = el
    holder._p.getparent().remove(holder._p)
    return d


# ——— Lekcja 31: sekcje, orientacja, kolumny, nagłówki ———
def lesson31():
    d = start(31)
    body = d.element.body
    harm = para(d, 'Harmonogram do sekcji poziomej')
    s1 = new_para_before(harm._p)                            # pusty akapit z podziałem sekcji (Następna strona)
    put_sect(s1, sect_xml())
    tbl = body.find(qn('w:tbl'))
    after = tbl.getnext()                                     # pusty akapit za tabelą
    put_sect(after, sect_xml(landscape=True))
    tips = [p for p in d.paragraphs if re.match(r'(Cel wydarzenia|Przygotowanie sali|Podział zadań|Przebieg turnieju|Zasady punktacji|Podsumowanie i wnioski)\. ', p.text)]
    s3 = new_para_before(tips[0]._p)                          # granica przed poradami (typ Ciągły ma sekcja 4)
    put_sect(s3, sect_xml())                             # sekcja 3 zaczyna się od nowej strony (po poziomej)
    put_sect(tips[-1]._p, sect_xml(cols=2, kind='continuous'))  # sectPr w ostatnim akapicie porad
    final = body.find(qn('w:sectPr'))
    final.insert(0, parse_xml(f'<w:type {W} w:val="continuous"/>'))
    secs = d.sections
    header_text(secs[0], 'PORADNIK ORGANIZATORA')
    header_text(secs[2], 'PORADNIK ORGANIZATORA')         # najpierw rozłączona sekcja ZA harmonogramem
    header_text(secs[1], 'PLAN WYDARZENIA')
    return d


# ——— Lekcja 32: recenzja i kopia publikacyjna ———
def ins(text, author='Uczeń Testowy'):
    return f'<w:ins w:id="{rid()}" w:author="{author}" w:date="{DATE}">{run_xml(text)}</w:ins>'


def dele(text, author='Uczeń Testowy'):
    return f'<w:del w:id="{rid()}" w:author="{author}" w:date="{DATE}"><w:r><w:delText xml:space="preserve">{esc(text)}</w:delText></w:r></w:del>'


def comment(doc, p, run_index, text):
    runs = [r for r in p._p.iter(qn('w:r')) if r.find(qn('w:t')) is not None]
    doc.add_comment(Run(runs[run_index], p), text=text, author='Uczeń Testowy', initials='UT')


def lesson32(variant='review'):
    d = start(32)
    p_place = para(d, 'Zapraszamy uczniów', exact=False)
    p_time = para(d, 'Warsztaty potrwają', exact=False)
    p_cert = para(d, 'Każdy uczestnik', exact=False)
    if variant == 'review':
        clear_runs(p_time)
        append_xml(p_time, run_xml('Warsztaty potrwają ') + dele('60 minut', 'Redaktor ćwiczeniowy') + ins('90 minut', 'Redaktor ćwiczeniowy')
                   + run_xml('. Udział jest ') + dele('płatny') + ins('bezpłatny') + run_xml('. Zgłoszenia przyjmuje opiekun koła do ')
                   + dele('17') + ins('16') + run_xml(' listopada.'))
        clear_runs(p_cert)
        append_xml(p_cert, dele('Każdy uczestnik otrzyma profesjonalny certyfikat. ') + run_xml('Na spotkaniu nauczysz ') + dele('sie') + ins('się')
                   + run_xml(' przygotowywać przejrzyste dokumenty. ') + run_xml('Przynieś materiały.'))
        comment(d, p_place, -2, 'Sala 12 — zgodna z ustaleniami. Akceptuję w wersji publicznej.')
        comment(d, p_time, 1, '90 minut jest niezgodne z kartą danych: ma zostać 60. Odrzucam.')
        comment(d, p_cert, -1, 'Do wyjaśnienia z organizatorem; do tego czasu usuńmy niepotwierdzoną prośbę z wersji publicznej.')
        comment(d, p_cert, 1, 'Po informacji zwrotnej poprawię zdanie o materiałach — usunę je z wersji publicznej.')
    else:  # kopia publikacyjna: zmiany rozstrzygnięte, komentarze usunięte, śledzenie wyłączone
        clear_runs(p_place)
        append_xml(p_place, run_xml('Zapraszamy uczniów klas pierwszych na warsztaty tworzenia dokumentów. Spotkanie odbędzie się 18 listopada 2026 o godzinie 14.00 w sali 12.'))
        clear_runs(p_time)
        append_xml(p_time, run_xml('Warsztaty potrwają 60 minut. Udział jest bezpłatny. Zgłoszenia przyjmuje opiekun koła do 16 listopada.'))
        clear_runs(p_cert)
        append_xml(p_cert, run_xml('Na spotkaniu nauczysz się przygotowywać przejrzyste dokumenty.'))
        tr = d.settings.element.find(qn('w:trackRevisions'))
        if tr is not None:
            d.settings.element.remove(tr)
    return d


# ——— Lekcja 33: poradnik ———
def lesson33(variant='final'):
    d = start(33)
    toc_styles(d)
    body = d.element.body
    if variant == 'review':
        p = para(d, 'Zasady punktacji', exact=False)
        text = para(d, 'Za zwycięstwo', exact=False)
        append_xml(text, ins(' Remis rozstrzygamy dodatkową rundą trwającą pięć minut.'))
        comment(d, text, 0, 'Kryterium: jasność zasad; problem: brak informacji o remisie; następny krok: dopisz rozstrzygnięcie.')
        comment(d, para(d, 'Stoły ustawiamy', exact=False), 0, 'Kryterium: konkret; problem: ile stołów?; następny krok: podaj liczbę.')
        return d
    para(d, 'Podsumowanie i wnioski').runs[0].text = 'Podsumowanie wydarzenia i wnioski'
    sala = para(d, 'Przygotowanie sali')
    new_para_after(para(d, 'Stoły ustawiamy', exact=False)._p, 'Heading2', run_xml('Lista zakupów'))
    new_para_after(para(d, 'Koordynator zbiera', exact=False)._p, 'Heading2', run_xml('Role w dniu turnieju'))
    # Podpisy
    pic = [p for p in d.paragraphs if p._p.findall('.//' + qn('w:drawing'))]
    hint1 = para(d, 'Dodaj automatyczny podpis z etykietą Ilustracja', exact=False)
    clear_runs(hint1)
    hint1.style = d.styles['Caption']
    append_xml(hint1, caption_xml('Ilustracja', 1, '. Model pracy zespołu: koordynator, sędzia i opiekun sprzętu', '_Toc4000001'))
    hint2 = para(d, 'Dodaj podpis z etykietą Wykres', exact=False)
    clear_runs(hint2)
    hint2.style = d.styles['Caption']
    append_xml(hint2, caption_xml('Wykres', 1, '. Najczęściej wskazywane potrzeby uczestników', '_Toc4000003'))
    tbl = body.find(qn('w:tbl'))
    cap_t = new_para_before(tbl, 'Caption', caption_xml('Tabela', 1, '. Harmonogram roboczy turnieju', '_Toc4000002'))
    after_hint = para(d, 'Tabela powinna znaleźć się', exact=False)
    clear_runs(after_hint)
    # Sekcja pozioma: od nagłówka „Harmonogram roboczy” do akapitu za tabelą
    harm = para(d, 'Harmonogram roboczy')
    put_sect(new_para_before(harm._p), sect_xml())
    put_sect(after_hint._p, sect_xml(landscape=True))
    # Spis treści po tytule
    entries = heading_entries(d)
    intro = para(d, 'Materiał startowy', exact=False)
    anchor = intro._p
    for el in toc_paragraphs(entries, 'TOC \\o "1-3" \\h \\z \\u', lambda lv: f'TOC{lv}', heading='Spis treści'):
        anchor.addnext(el)
        anchor = el
    # Spisy obiektów przed „Źródła materiałów”
    src = para(d, 'Źródła materiałów')._p
    for label, title, bm, text in [('Ilustracja', 'Spis ilustracji', '_Toc4000001', 'Ilustracja 1. Model pracy zespołu: koordynator, sędzia i opiekun sprzętu'),
                                   ('Tabela', 'Spis tabel', '_Toc4000002', 'Tabela 1. Harmonogram roboczy turnieju'),
                                   ('Wykres', 'Spis wykresów', '_Toc4000003', 'Wykres 1. Najczęściej wskazywane potrzeby uczestników')]:
        new_para_before(src, None, f'<w:r><w:rPr><w:b/></w:rPr><w:t>{title}</w:t></w:r>')
        for el in toc_paragraphs([(text, 1, bm, 4)], f'TOC \\h \\z \\c "{label}"', lambda lv: 'TableofFigures', sdt=False):
            src.addprevious(el)
    secs = d.sections
    header_text(secs[0], 'PORADNIK ORGANIZATORA')
    footer_page(secs[0], '', 'sdt')
    tr = d.settings.element.find(qn('w:trackRevisions'))
    if tr is not None:
        d.settings.element.remove(tr)
    return d


PL_STYLES = [('Heading1', 'Nagwek1', 'Nagłówek 1'), ('Heading2', 'Nagwek2', 'Nagłówek 2'), ('Title', 'Tytu', 'Tytuł'), ('Normal', 'Normalny', 'Normalny')]

if __name__ == '__main__':
    save(lesson28(), '28-solved.docx')
    save(lesson28('pl'), '28-solved-pl.docx', stored=True, rename_styles=PL_STYLES)
    save(lesson28('fake'), '28-fake-heading.docx')
    save(lesson28('manual'), '28-manual-page.docx')
    save(lesson29(), '29-solved.docx')
    save(lesson29('stale'), '29-stale-toc.docx')
    save(lesson30(), '30-solved.docx')
    save(lesson31(), '31-solved.docx')
    save(lesson32('review'), '32-recenzja.docx')
    save(lesson32('publish'), '32-publikacja.docx')
    save(lesson33('final'), '33-solved.docx')
    save(lesson33('review'), '33-recenzja.docx')
    ok = (OUT / '28-solved.docx').read_bytes()
    (OUT / 'broken.docx').write_bytes(ok[: len(ok) * 2 // 3])
    print('broken.docx: obcięty ZIP')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w') as z:
        z.writestr('[Content_Types].xml', '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        z.writestr('xl/workbook.xml', '<workbook/>')
    (OUT / 'not-word.xlsx').write_bytes(buf.getvalue())
    print('not-word.xlsx: archiwum Excela')
