#!/usr/bin/env python3
"""Build independent Word desktop classroom starters. Requires python-docx, Pillow and lxml."""
from pathlib import Path
import json, zipfile
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from lxml import etree
from PIL import Image, ImageDraw, ImageFont
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'public/materials/word'; ASSETS=OUT/'assets'; ASSETS.mkdir(parents=True,exist_ok=True)
FONT='/System/Library/Fonts/Supplemental/Arial.ttf'
font=ImageFont.truetype(FONT,28);small=ImageFont.truetype(FONT,23)
for name,title,labels,vals in [('chart-usage','Czas korzystania ze strefy nauki',['Pon.','Wt.','Śr.','Czw.','Pt.'],[32,45,38,52,43]),('chart-survey','Potrzeby uczestników',['Cisza','Gniazdka','Stół grupowy','Tablica'],[18,15,12,9])]:
 im=Image.new('RGB',(1200,430),'white');dr=ImageDraw.Draw(im);dr.text((45,20),title,fill='black',font=font)
 dr.text((45,65),'Liczba wskazań' if 'survey' in name else 'Minuty',fill='#45555a',font=small)
 for i,(lab,val) in enumerate(zip(labels,vals)):
  x=130+i*220;y=360-val/max(vals)*225
  dr.rectangle((x,y,x+100,360),fill='#347c90');dr.text((x+35,y-32),str(val),font=small,fill='black');dr.text((x-10,375),lab,font=small,fill='black')
 im.save(ASSETS/(name+'.png'))
for name,title,boxes in [('illustration-plan','Plan strefy nauki',['Cicha praca','Wspólny stół','Regał materiałów']),('illustration-process','Droga od pomysłu do gotowej pracy',['Zaplanuj','Wykonaj','Sprawdź'])]:
 im=Image.new('RGB',(1200,300),'white');dr=ImageDraw.Draw(im);dr.text((45,20),title,fill='black',font=font)
 for i,t in enumerate(boxes):
  x=35+i*400;dr.rounded_rectangle((x,105,x+325,235),radius=16,fill='#eaf3f5',outline='#347c90',width=3);dr.text((x+162,170),t,font=font,fill='black',anchor='mm')
  if i<2:dr.line((x+340,170,x+380,170),fill='#347c90',width=4);dr.polygon([(x+380,170),(x+370,163),(x+370,177)],fill='#347c90')
 im.save(ASSETS/(name+'.png'))
def base(title,intro):
 d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=s.bottom_margin=Cm(2);s.left_margin=s.right_margin=Cm(2.2)
 for n in ['Normal','Title','Heading 1','Heading 2']:
  st=d.styles[n];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
 for border in list(d.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.space_after=Pt(8)
 d.styles['Normal'].paragraph_format.line_spacing=1.12
 d.styles['Title'].font.size=Pt(25);d.styles['Heading 1'].font.size=Pt(17);d.styles['Heading 2'].font.size=Pt(13)
 d.core_properties.author='Pracownia informatyczna';d.core_properties.title=title
 d.add_paragraph(title,'Title');d.add_paragraph(intro);return d
def p(d,t):return d.add_paragraph(t)
def h(d,t,l=1):return d.add_paragraph(t,'Heading '+str(l))
def table(d,heads,rows):
 t=d.add_table(rows=1,cols=len(heads));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,x in zip(t.rows[0].cells,heads):c.text=x
 for row in rows:
  for c,x in zip(t.add_row().cells,row):c.text=str(x)
 for i,r in enumerate(t.rows):
  for c in r.cells:
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   pr.append(b);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCECF0' if i==0 else ('F5F7F8' if i%2==0 else 'FFFFFF'));pr.append(sh)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(5);para.paragraph_format.space_before=Pt(5)
    for run in para.runs:run.font.size=Pt(10);run.bold=i==0
  if i==0: r._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 p(d,'');return t
def pic(d,name,alt):
 sh=d.add_picture(str(ASSETS/(name+'.png')),width=Cm(15));sh._inline.docPr.set('descr',alt)
def save(d,n):d.save(OUT/f'lesson-{n}-start.docx')
def page(d):d.add_page_break()
# 28: purposefully plain section labels for applying semantic styles.
d=base('Szkolny klub gier','Dokument roboczy zespołu uczniowskiego. Przekształć go w czytelny informator klubu. Tytuły działów i podrozdziałów mają teraz styl Normalny. Dodaj style, nagłówek i stopkę w aplikacji Word na komputerze.')
for title,paras in [
 ('Po co zakładamy klub',['Klub jest miejscem wspólnego poznawania gier planszowych i logicznych. Uczestnicy ćwiczą wyjaśnianie zasad, planowanie działań oraz uczciwe rozwiązywanie sporów. Nie trzeba mieć własnej gry ani doświadczenia.','Spotkania odbywają się w szkolnej sali po lekcjach. Przed każdym spotkaniem zespół wybiera grę odpowiednią do dostępnego czasu i liczby osób. Organizator sprawdza kompletność elementów.']),
 ('Pierwsze spotkanie',['Najpierw uczestnicy przedstawiają swoje oczekiwania. Następnie poznają krótką grę i rozgrywają rundę próbną. Po próbie można jeszcze raz wyjaśnić trudniejsze reguły.']),
 ('Role w zespole',['Osoba prowadząca tłumaczy zasady. Opiekun materiałów wydaje pudełka i sprawdza ich zawartość. Reporter zapisuje wnioski, a gospodarz sali pilnuje czasu i porządku. Role zmieniamy co spotkanie.']),
 ('Zasady współpracy',['Pozwalamy innym dokończyć wypowiedź. Spór o regułę rozwiązujemy przez sprawdzenie instrukcji. Nie zmieniamy zasad w połowie partii bez zgody wszystkich uczestników.','Wygrana nie jest jedynym celem. Po rozgrywce każdy wskazuje jedną decyzję, z której jest zadowolony, i jedną rzecz, którą następnym razem zrobi inaczej.']),
 ('Co przygotować',['Potrzebujemy dwóch kompletnych gier, zegara, arkusza zapisów i miejsca na odłożenie plecaków. Materiały przechowujemy w opisanych pudełkach. Nie zapisujemy danych osobowych uczestników w ogólnodostępnym pliku.']),
 ('Jak sprawdzimy efekty',['Na koniec spotkania zbieramy krótkie odpowiedzi: co było zrozumiałe, co wymaga wyjaśnienia i którą grę warto poznać następnym razem. Na ich podstawie zespół planuje następne zajęcia.'])]:
 p(d,title).runs[0].bold=True
 for txt in paras:p(d,txt)
 if title=='Jak sprawdzimy efekty':
  for extra in ['Reporter zapisuje trzy najczęściej pojawiające się uwagi. Każda powinna dotyczyć zachowania lub rozwiązania organizacyjnego, a nie cech konkretnej osoby. Zamiast pisać, że ktoś był słaby, opisujemy moment, w którym instrukcja okazała się niejasna.', 'Na kolejnym spotkaniu wracamy do wybranego wniosku. Jeśli poprzednio brakowało czasu na rundę próbną, planujemy krótsze powitanie. Po zmianie sprawdzamy, czy uczestnicy lepiej rozumieją reguły.', 'Informator klubu aktualizujemy po wspólnej decyzji zespołu. W nagłówku umieść krótką nazwę klubu, a w stopce automatyczny numer strony. Po zmianie stylu Normalny sprawdź cały dokument i upewnij się, że nadal jest czytelny.']:p(d,extra)
 if title in ['Role w zespole','Co przygotować']:page(d)
save(d,28)
# 29
DRAFTS=[('Cel wydarzenia','Turniej ma umożliwić uczniom poznanie nowych osób i sprawdzenie różnych strategii gry. Wybieramy krótkie rozgrywki, aby każdy mógł uczestniczyć w kilku rundach.'),('Przygotowanie sali','Stoły ustawiamy tak, aby pomiędzy nimi było swobodne przejście. Każde stanowisko otrzymuje numer. Osobny stolik służy do zapisów i przekazywania wyników.'),('Podział zadań','Koordynator zbiera zgłoszenia, sędzia wyjaśnia reguły, a opiekun sprzętu przygotowuje materiały. Każdy zna osobę, do której zgłasza problem.'),('Przebieg turnieju','Po powitaniu przedstawiamy harmonogram i zasady. Runda próbna pozwala zauważyć niejasności. Wyniki zapisujemy po każdej rundzie, a uczestnicy potwierdzają ich poprawność.'),('Zasady punktacji','Za zwycięstwo przyznajemy trzy punkty, za remis jeden, a za przegraną zero. W razie równej liczby punktów stosujemy wcześniej ogłoszoną zasadę rozstrzygnięcia.'),('Podsumowanie i wnioski','Zbieramy opinie o długości rund i czytelności zasad. Wyniki omawiamy bez oceniania osób. Kolejna edycja korzysta z pomysłów zapisanych w podsumowaniu.')]
d=base('Organizacja szkolnego turnieju','Tekst do ćwiczenia automatycznego spisu treści. Główne działy mają styl Nagłówek 1, podrozdziały Nagłówek 2. Wstaw automatyczny spis pod tym wstępem, a następnie przetestuj jego aktualizację.')
p(d,'Miejsce na automatyczny spis treści');page(d)
for i,(title,txt) in enumerate(DRAFTS):
 h(d,title);p(d,txt);h(d,'Zadanie zespołu',2);p(d,'Ustalcie konkretne rozwiązanie, osobę odpowiedzialną i termin sprawdzenia. Zapiszcie decyzję pełnym zdaniem. W razie zmiany warunków wróćcie do tego ustalenia i uaktualnijcie opis.')
 if i==2:page(d)
save(d,29)
#30
d=base('Strefa nauki w naszej szkole','Raport ćwiczeniowy zawiera dwa rysunki, dwie tabele i dwa wykresy. Wszystkie dane są przykładowe. Elementy nie mają jeszcze podpisów Worda. Dodaj podpisy z trzema osobnymi etykietami Ilustracja, Tabela i Wykres oraz trzy automatyczne spisy.')
h(d,'Projekt przestrzeni');p(d,'Plan przedstawia trzy miejsca o różnych funkcjach. Podpisz rysunek jako Ilustracja i dopisz zwięzły opis.');pic(d,'illustration-plan','Schemat trzech miejsc: cicha praca, wspólny stół, regał z materiałami.')
p(d,'Drugi rysunek przedstawia etapy pracy. Także on należy do kategorii Ilustracja.');pic(d,'illustration-process','Trzy kolejne etapy: zaplanuj, wykonaj, sprawdź.');page(d)
h(d,'Wyposażenie i organizacja');p(d,'Obie poniższe tabele wymagają podpisów z etykietą Tabela.')
table(d,['Element','Liczba','Przeznaczenie'],[['Stolik',4,'Praca indywidualna'],['Duży stół',1,'Wspólny projekt'],['Krzesło',12,'Miejsca do pracy'],['Tablica',1,'Zapis pomysłów']])
table(d,['Rola','Zadanie','Termin'],[['Gospodarz','Sprawdzenie porządku','Przed otwarciem'],['Dyżurny','Wydawanie materiałów','Podczas spotkania'],['Reporter','Zapis wniosków','Po spotkaniu']])
p(d,'W danych nie ma informacji o rzeczywistej szkole ani o konkretnych uczniach. Możesz porównać tabele, ale nie zmieniaj danych liczbowych bez uzasadnienia.');page(d)
h(d,'Obserwacje zespołu');p(d,'Pierwszy wykres pokazuje średni czas jednej wizyty w kolejnych dniach. Drugi przedstawia wskazania w przykładowej ankiecie wielokrotnego wyboru. Podpisz je etykietą Wykres. Wykresy są obrazami, lecz automatyczny podpis działa przy nich tak samo.');pic(d,'chart-usage','Minuty w kolejnych dniach: 32, 45, 38, 52, 43.');pic(d,'chart-survey','Potrzeby: cisza 18, gniazdka 15, stół grupowy 12, tablica 9.');page(d)
h(d,'Spisy elementów');p(d,'Tutaj wstaw trzy osobne spisy automatyczne. W oknie Wstaw spis ilustracji za każdym razem wybierz inną etykietę podpisu.');save(d,30)
#31
d=base('Plan szkolnego wydarzenia','Dokument do ćwiczenia podziałów strony i sekcji. Cały plik ma początkowo jedną sekcję pionową. Utwórz oddzielną poziomą sekcję dla tabeli i kolejną pionową dla porad. Same porady podziel na dwie kolumny.')
h(d,'Założenia wydarzenia');p(d,'Organizujemy popołudnie z grami i łamigłówkami. Uczestnicy wybierają stanowiska i pracują w małych grupach. Wydarzenie ma pomóc w poznaniu osób z innych klas oraz pokazać różne sposoby uczenia się.')
p(d,'Otwarcie zawiera powitanie i wyjaśnienie reguł. Część główna obejmuje trzy rundy. Na koniec uczestnicy porządkują stanowiska i przekazują po jednej propozycji usprawnienia. W planie pozostawiamy zapas czasu na zmianę grup.')
h(d,'Harmonogram do sekcji poziomej');p(d,'Wstaw podział sekcji Następna strona przed tym nagłówkiem i po tabeli. Tylko tę sekcję ustaw poziomo; dopasuj tabelę do szerokości strony.')
table(d,['Godzina','Etap','Miejsce','Osoba','Materiały','Wynik'],[['14.00','Powitanie','Sala A','Prowadzący','Plan dnia','Znane zasady'],['14.10','Runda 1','Stoły 1–4','Opiekunowie','Gry logiczne','Pierwsza próba'],['14.30','Zmiana','Sala A','Dyżurny','Numery grup','Gotowość'],['14.35','Runda 2','Stoły 1–4','Opiekunowie','Łamigłówki','Nowa strategia'],['14.55','Runda 3','Stoły 1–4','Opiekunowie','Karty zadań','Współpraca'],['15.15','Podsumowanie','Sala A','Reporter','Kartki','Wnioski']])
h(d,'Porady dla organizatorów');p(d,'Od tego nagłówka dokument ma ponownie układ pionowy. Sam tekst porad umieść w dwóch kolumnach, korzystając w razie potrzeby z podziału sekcji Ciągły.')
for title,txt in DRAFTS:
 p(d,title+'. '+txt+' Sprawdź ustalenie przed wydarzeniem i przekaż je pozostałym osobom w zespole. Jeżeli pojawi się trudność, opisz ją konkretnie i zaproponuj możliwe rozwiązanie.')
h(d,'Zamknięcie dokumentu');p(d,'Ten akapit ma być znów jednokolumnowy. Użyj podziału sekcji Ciągły po poradach. Włącz znaki niedrukowane i sprawdź położenie wszystkich podziałów.');save(d,31)
#32
d=base('Zaproszenie na warsztaty','Pakiet recenzyjny. Pierwsza część zawiera uzgodnione dane, a druga tekst do korekty. Korzystaj z Recenzja i Wszystkie adiustacje. Plik zawiera już dwie propozycje zmian; każdą rozpatrz oddzielnie.')
h(d,'Uzgodnione dane do sprawdzenia');table(d,['Pole','Obowiązujące ustalenie'],[['Termin','18 listopada 2026, godzina 14.00'],['Miejsce','Sala 12'],['Czas','60 minut'],['Koszt','Udział bezpłatny'],['Zapisy','Do 16 listopada u opiekuna koła'],['Uczestnicy','Uczniowie klas pierwszych']])
p(d,'Zasady redakcji: pisz uprzejmie i konkretnie. Nie dodawaj obietnic niepotwierdzonych w ustaleniach. Popraw błędy z włączonym śledzeniem zmian. Zostaw komentarz tam, gdzie brakuje informacji, zamiast ją wymyślać. Dane opisują fikcyjne wydarzenie ćwiczeniowe.')
h(d,'Tekst zaproszenia do recenzji');p(d,'Zapraszamy uczniów klas pierwszych na warsztaty tworzenia dokumentów. Spotkanie odbędzie się 18 listopada 2026 o godzinie 14.00 w sali 21.');p(d,'Warsztaty potrwają 60 minut. Udział jest płatny. Zgłoszenia przyjmuje opiekun koła do 17 listopada.');p(d,'Każdy uczestnik otrzyma profesjonalny certyfikat. Na spotkaniu nauczysz sie przygotowywać przejrzyste dokumenty. Przynieś materiały.');p(d,'Serdecznie zapraszamy do wspólnej pracy. Pytania możesz przekazać opiekunowi koła.');save(d,32)
# Add tracked replacements in exactly two runs, preserving real accepting/rejecting workflow.
path=OUT/'lesson-32-start.docx'
with zipfile.ZipFile(path) as z: members={n:z.read(n) for n in z.namelist()}
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'};root=etree.fromstring(members['word/document.xml']);cid=1
for old,new in [('sali 21','sali 12'),('60 minut','90 minut')]:
 # target invitation only: last occurrence, never edit source brief
 node=[n for n in root.findall('.//w:t',ns) if old in (n.text or '')][-1];run=node.getparent();parent=run.getparent();idx=parent.index(run);a,b=node.text.split(old,1);repl=[]
 for kind,txt in [('r',a),('del',old),('ins',new),('r',b)]:
  if not txt:continue
  e=etree.Element(qn('w:'+kind))
  if kind=='r':rr=e
  else:
   e.set(qn('w:id'),str(cid));cid+=1;e.set(qn('w:author'),'Redaktor ćwiczeniowy');e.set(qn('w:date'),'2026-09-30T08:00:00Z');rr=etree.SubElement(e,qn('w:r'))
  tt=etree.SubElement(rr,qn('w:delText' if kind=='del' else 'w:t'));tt.set('{http://www.w3.org/XML/1998/namespace}space','preserve');tt.text=txt;repl.append(e)
 parent.remove(run)
 for off,e in enumerate(repl):parent.insert(idx+off,e)
settings=etree.fromstring(members['word/settings.xml']);settings.append(etree.Element(qn('w:trackRevisions')))
members['word/document.xml']=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True);members['word/settings.xml']=etree.tostring(settings,xml_declaration=True,encoding='UTF-8',standalone=True)
with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
 for n,b in members.items():z.writestr(n,b)
#33
d=base('Poradnik organizacji szkolnego wydarzenia','Materiał startowy do samodzielnego opracowania poradnika dla uczniowskiego zespołu. Zredaguj tekst i przygotuj 4–6 czytelnych stron. Połącz style, spis treści, podpisy i spisy obiektów, sekcje oraz recenzję. Informacje dotyczą przykładowego wydarzenia.')
p(d,'Odbiorca: uczeń, który organizuje wydarzenie po raz pierwszy. Cel: po przeczytaniu poradnika potrafi zaplanować zadania, przygotować salę i zebrać wnioski. Uzupełnij własny tytuł, autora oraz opis wydarzenia.')
for i,(title,txt) in enumerate(DRAFTS):
 h(d,title);p(d,txt);p(d,'W praktyce zapisz ustalenia tak, aby mogła z nich skorzystać osoba nieobecna podczas rozmowy. Podaj kolejność działań i sposób sprawdzenia wyniku. Unikaj sformułowań „zrób to dobrze” bez wyjaśnienia, co dokładnie należy zrobić.')
 if i==1:
  h(d,'Model pracy zespołu',2);pic(d,'illustration-process','Etapy przygotowania: zaplanuj, wykonaj, sprawdź.');p(d,'Dodaj automatyczny podpis z etykietą Ilustracja. Opisz, jak model pracy pomaga twojemu zespołowi.')
 if i==3:
  h(d,'Harmonogram roboczy',2);table(d,['Etap','Czas','Odpowiedzialność','Zasoby','Sprawdzenie'],[['Otwarcie','10 min','Prowadzący','Plan','Zasady znane'],['Warsztaty','40 min','Zespół','Materiały','Zadanie gotowe'],['Podsumowanie','10 min','Reporter','Ankieta','Wnioski zapisane']]);p(d,'Tabela powinna znaleźć się w osobnej sekcji poziomej. Dodaj podpis z etykietą Tabela, następnie wróć do układu pionowego.')
h(d,'Jak czytać zebrane opinie');p(d,'Przykładowe wskazania uczestników pokazują, które elementy miejsca pracy są ważne. To dane treningowe, a nie wyniki ankiety w twojej szkole. Nie traktuj wskazań jako procentów, bo można było wybrać kilka odpowiedzi.');pic(d,'chart-survey','Cisza 18, gniazdka 15, stół grupowy 12, tablica 9.');p(d,'Dodaj podpis z etykietą Wykres. Napisz dwa wnioski zgodne z danymi i jedną informację, której nie da się z nich ustalić.')
h(d,'Lista kontroli przed publikacją');p(d,'Sprawdź hierarchię nagłówków w okienku nawigacji. Zaktualizuj spis treści oraz spisy ilustracji, tabel i wykresów. Skontroluj numerację stron oraz powrót do pionowej orientacji po tabeli.');p(d,'Poproś partnera o dwie konkretne uwagi w komentarzach i jedną propozycję zmiany ze śledzeniem. Podejmij decyzję o każdej zmianie. Zapisz wersję roboczą z recenzją i osobną wersję finalną po rozpatrzeniu uwag.');h(d,'Źródła materiałów');p(d,'Tekst, schematy i dane ćwiczeniowe przygotowano na potrzeby tych zajęć. Jeśli dodasz materiał z internetu, podaj autora, tytuł, adres i warunki użycia. Nie kopiuj cudzych zdjęć bez uprawnienia.');save(d,33)
manifest={'version':1,'target':'Microsoft Word na komputerze, domyślnie Windows','fixtures':[
 {'lesson':28,'filename':'lesson-28-start.docx','topic':'Szkolny klub gier','expectedOperations':['Apply Heading 1/2 styles','Insert header and PAGE footer field','Modify Normal style'],'initialSections':1},
 {'lesson':29,'filename':'lesson-29-start.docx','topic':'Organizacja szkolnego turnieju','expectedOperations':['Insert TOC field','Add or rename heading','Update entire TOC'],'prestyledHeadings':True,'initialTocCount':0},
 {'lesson':30,'filename':'lesson-30-start.docx','topic':'Strefa nauki','expectedOperations':['Add 2 Ilustracja captions','Add 2 Tabela captions','Add 2 Wykres captions','Insert 3 separate label-filtered TOCs'],'illustrations':2,'tables':2,'chartImages':2,'initialCaptionFields':0},
 {'lesson':31,'filename':'lesson-31-start.docx','topic':'Plan szkolnego wydarzenia','expectedOperations':['Insert next-page section breaks around schedule','Set schedule landscape','Return to portrait','Use continuous section breaks for two-column tips'],'initialSections':1,'initialColumns':1},
 {'lesson':32,'filename':'lesson-32-start.docx','topic':'Zaproszenie na warsztaty','expectedOperations':['Accept room correction 21 to 12','Reject duration correction 60 to 90','Track own factual and language corrections','Add comment for missing materials detail'],'trackedInsertions':2,'trackedDeletions':2,'truth':{'date':'18 listopada 2026','time':'14.00','room':'12','minutes':60,'cost':'bezpłatny','deadline':'16 listopada'},'initialComments':0},
 {'lesson':33,'filename':'lesson-33-start.docx','topic':'Poradnik organizacji szkolnego wydarzenia','expectedOperations':['Create 4–6 page guide','Apply heading hierarchy','Insert header and PAGE footer field','Insert TOC and caption lists','Landscape table section','Peer review with tracked change and comments'],'illustrations':1,'tables':1,'chartImages':1}]}
(OUT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
print('Built six DOCX starters and manifest in',OUT)
