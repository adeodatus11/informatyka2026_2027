#!/usr/bin/env python3
# Generator pliku startowego lekcji 42 (stanowisko trójki, 1 strona). Wymaga python-docx.
# Użycie: python3 scripts/build-lesson-42-docx.py public/materials/lesson-42-stanowisko-start.docx
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

out=sys.argv[1]
d=Document()
sec=d.sections[0]
sec.page_height=Cm(29.7);sec.page_width=Cm(21.0)
for m in ('left_margin','right_margin'):setattr(sec,m,Cm(2.0))
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8)
st=d.styles['Normal'];st.font.name='Calibri';st.font.size=Pt(11)
st.element.rPr.rFonts.set(qn('w:eastAsia'),'Calibri')
st.paragraph_format.space_after=Pt(3);st.paragraph_format.line_spacing=1.1
d.core_properties.title='Stanowisko trójki — szanse i zagrożenia technologii'
d.core_properties.author='Informatyka praktycznie — materiał ćwiczeniowy'

def para(text,bold=False,italic=False,size=None,color=None,style=None,after=None):
    p=d.add_paragraph(style=style)
    r=p.add_run(text);r.bold=bold;r.italic=italic
    if size:r.font.size=Pt(size)
    if color:r.font.color.rgb=RGBColor.from_string(color)
    if after is not None:p.paragraph_format.space_after=Pt(after)
    return p
def section(title):
    # Celowo: wygląda jak nagłówek, ale to styl Normalny — uczeń nadaje Nagłówek 1.
    p=para(title,bold=True,size=13,color='1F4FA3',after=2)
    p.paragraph_format.space_before=Pt(8);p.paragraph_format.keep_with_next=True
    return p

t=d.add_paragraph('[Tytuł Waszego stanowiska]',style='Title')
para('Temat debaty: AI i nowe technologie w szkole i w pracy — więcej szans czy zagrożeń?',italic=True,color='4A5A6B',after=4)
section('Nasza teza')
para('[Wklejcie tezę z kreatora: „Więcej szans…” albo „Więcej zagrożeń…”]')
section('Argumenty')
para('Argument 1 ([rola]): [treść argumentu]')
para('Dowód: [kto, kiedy, co zbadał — liczby z raportu oznaczcie „do weryfikacji”]',color='4A5A6B',after=6)
para('Argument 2 ([rola]): [treść argumentu]')
para('Dowód: [kto, kiedy, co zbadał]',color='4A5A6B')
section('Druga strona mówi')
para('[Najmocniejszy argument drugiej strony]')
para('Dowód: [źródło]',color='4A5A6B')
section('Nasza odpowiedź')
para('[Co jest trafne i jaką zasadę proponujecie]')
section('Szanse i zagrożenia w skrócie')
tb=d.add_table(rows=3,cols=2);tb.style='Table Grid'
for i,h in enumerate(['Szanse','Zagrożenia']):
    c=tb.rows[0].cells[i];c.text='';r=c.paragraphs[0].add_run(h);r.bold=True
    sh=OxmlElement('w:shd');sh.set(qn('w:val'),'clear');sh.set(qn('w:fill'),'EEF3FC');c._tc.get_or_add_tcPr().append(sh)
for row in tb.rows[1:]:
    for c in row.cells:c.text='[hasło, 3–6 słów]'
section('Nasza zasada')
para('[Kto, co i kiedy robi — tak, żeby dało się sprawdzić]')
section('Dbaj o siebie')
para('[Jedna wskazówka dobrostanu cyfrowego]')
para('Materiał ćwiczeniowy. Dane z raportów sprawdźcie u źródła przed publikacją. Autorów i licencję wpiszcie w stopce strony.',italic=True,size=9,color='4A5A6B').paragraph_format.space_before=Pt(10)
d.save(out)
